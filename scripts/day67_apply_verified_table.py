"""等价表 v1 构建：读用户《等价表核验·最终总表.xlsx》→ data/jobs/equivalence_table.json。

用户核验结论（2026-10-23，xlsx 汇总说明）：
- 高频 163：✓ 138｜✗ 2（#64/#76 删除）｜A组拆解 14（复合→原子，✓半舍半）｜D组 1（#10 ✓·减分）
- B组 3（工具名≠能力）+ C组 5（角色≠方法）不勾 → 不入表
- 低频 93（原采纳 31 + 本轮扩充 62）全部采纳

表结构（coverage_engine 消费）：
- equivalences: {要求侧词: [证据侧词...]}——同词对去重、norm 后写入
- deduct: [[k,v]]——D组 #10 受众错位，命中按 0.8 计（用户「酌情减分」决议）
- atomic_splits: {复合要求词: [原子能力...]}——score_coverage 把含该复合的要求
  拆成等权子要求；匹配半边经 equivalences 命中、舍弃半边无表项保持未证实（防半匹配污染）

用法：python scripts/day67_apply_verified_table.py [--xlsx 路径]
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from utils.coverage_engine import _norm_literal  # noqa: E402

TABLE_OUT = ROOT / "data" / "jobs" / "equivalence_table.json"
DEFAULT_XLSX = r"D:\xwechat_files\wxid_ssjx8zam72wp22_6a40\msg\file\2026-10\等价表核验·最终总表.xlsx"


def _clean(s) -> str:
    return str(s or "").strip() if s is not None else ""


def build(xlsx_path: str) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    equiv: dict[str, list[str]] = {}
    deduct: list[list[str]] = []
    splits: dict[str, list[str]] = {}
    stats = {"high_ok": 0, "high_x": 0, "atomic_ok": 0, "atomic_drop": 0,
             "b_c_skip": 0, "low_ok": 0, "deduct": 0}

    def add(k: str, v: str) -> None:
        kn, vn = _norm_literal(k), _norm_literal(v)
        if not kn or not vn or kn == vn:
            return
        bucket = equiv.setdefault(kn, [])
        if vn not in bucket:
            bucket.append(vn)

    # ── 高频 163 ──
    for row in wb["高频163条核验"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        num, req, chunk, _rel, _cnt, _emps, _jobs, ver, _note = (list(row) + [None] * 9)[:9]
        v = _clean(ver)
        if v == "✓":
            add(req, chunk)
            stats["high_ok"] += 1
        elif v == "✓·减分":
            add(req, chunk)
            deduct.append([_norm_literal(req), _norm_literal(chunk)])
            stats["deduct"] += 1
        elif v == "✗":
            stats["high_x"] += 1  # 用户否决：不入表
        elif v.startswith("拆解"):
            stats["b_c_skip"] += 0  # 拆解行由 A组 sheet 承载
        elif v.startswith("不勾"):
            stats["b_c_skip"] += 1  # B/C 组：不入表
        # 其余（空）不勾

    # ── A组拆解原子条目：✓ 行入表（键=原子词）；拆解组登记 atomic_splits ──
    for row in wb["A组拆解原子条目"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        _src, compound, atomic, chunk, _rel, verdict, _note = (list(row) + [None] * 7)[:7]
        cn = _norm_literal(compound)
        splits.setdefault(cn, [])
        an = _norm_literal(atomic)
        if an and an not in splits[cn]:
            splits[cn].append(an)
        if _clean(verdict) == "✓":
            add(atomic, chunk)
            stats["atomic_ok"] += 1
        else:
            stats["atomic_drop"] += 1  # 舍弃半边：无表项，保持未证实

    # ── 低频 93：全部采纳 ──
    for row in wb["低频采纳总表"].iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        _num, req, chunk, _rel, _cnt, ver, _src = (list(row) + [None] * 7)[:7]
        if _clean(ver) == "✓":
            add(req, chunk)
            stats["low_ok"] += 1

    return {
        "version": "v1_20261023",
        "created": datetime.now().isoformat(timespec="seconds"),
        "provenance": (
            "用户《等价表核验·最终总表.xlsx》人工核验（2026-10-23）：高频✓138 + 拆解原子✓14"
            f" + D组减分1 + 低频✓93 = {stats['high_ok'] + stats['atomic_ok'] + stats['deduct'] + stats['low_ok']} 条映射；"
            "✗删除2（#64 方案设计≠项目管理 / #76 主动学习前沿=成长型思维词根）；"
            "B组3（工具名≠能力）C组5（角色≠方法）不勾不入表；"
            "atomic_splits=复合要求拆原子（匹配半边入表、舍弃半边保持未证实，防半匹配污染）；"
            "deduct=受众错位条目按 0.8 计（用户酌情减分决议）。"
            "上游：Round3 §三 全量 11,920 对 LLM 离线提议 → 人工核验；运行时只查表零 LLM。"
        ),
        "stats": stats,
        "equivalences": equiv,
        "deduct": deduct,
        "atomic_splits": splits,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    args = ap.parse_args()
    table = build(args.xlsx)
    TABLE_OUT.write_text(json.dumps(table, ensure_ascii=False, indent=1), encoding="utf-8")
    n_pairs = sum(len(v) for v in table["equivalences"].values())
    print(f"等价表 {table['version']} → {TABLE_OUT.name}")
    print(f"  键 {len(table['equivalences'])} 个｜映射对 {n_pairs} 条｜减分对 {len(table['deduct'])}"
          f"｜原子拆解组 {len(table['atomic_splits'])}")
    print(f"  统计 {table['stats']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
